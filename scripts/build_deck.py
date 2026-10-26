"""content/<강>/slides.yaml → build/pptx/<번호>_<제목>.pptx

사용: python scripts/build_deck.py [강 번호 ...]   (생략 시 content/ 에 있는 강 전체)
"""
import hashlib
import sys
from pathlib import Path

import qrcode
import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
THEME = yaml.safe_load((ROOT / "config" / "theme.yaml").read_text(encoding="utf-8"))
COURSE = yaml.safe_load((ROOT / "config" / "course.yaml").read_text(encoding="utf-8"))
FONT = THEME["font"]
COL = THEME["colors"]
SOFT_TEXT = "D8DBEA"          # 어두운 배경 위 본문
CALLOUT_FILL = "FFF4D6"       # amber 옅은 바탕

W, H = 13.333, 7.5
M = 0.6                        # 바깥 여백
missing = []


# ── 기본 도형 ─────────────────────────────────────────────

def rgb(c):
    return RGBColor.from_string(COL.get(c, c))


def style_run(run, size, color, bold=False):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, FONT
    f.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ko-KR")        # 한국어로 지정해야 단어 단위로 줄바꿈된다
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", FONT)        # 한글 글꼴 지정


def text(slide, x, y, w, h, body, size, color="ink", bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(body).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing = align, spacing
        p._p.get_or_add_pPr().set("latinLnBrk", "0")  # "한글 단어 잘림 허용" 끄기
        r = p.add_run()
        style_run(r, size, color, bold)
        r.text = line
    return box


def shape(slide, kind, x, y, w, h, fill=None, line=None, radius=None, dash=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.5)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius / min(w, h)
    return s


def badge(slide, x, y, d, fill, label, size=18, color="paper"):
    """색 원 + 가운데 글자 (번호·O/X 표시)."""
    c = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    style_run(r, size, color, bold=True)
    r.text = label
    return c


def background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(color)


# ── 그림 ─────────────────────────────────────────────────

def illustration(slide, ref, x, y, w, h, dark):
    """assets/illustrations/<ref>.png 를 꽉 채워(잘라서) 둥근 액자로. 없으면 자리표시."""
    path = next((p for ext in (".png", ".jpg", ".webp")
                 if (p := ROOT / "assets" / "illustrations" / f"{ref}{ext}").exists()), None)
    if path is None:
        missing.append(ref)
        ph = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h,
                   fill="stage_2" if dark else "soft", line="muted", radius=0.25, dash=True)
        ph.name = f"자리표시 {ref}"
        text(slide, x, y, w, h,
             f"Google Flow 이미지 자리\n{ref}\nflow/prompts/{ref.split('/')[0]}_lecture.md",
             14, "muted", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.4)
        return
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:                       # 이미지가 더 넓다 → 좌우 자르기
        cut = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = pic.crop_right = cut
    else:                                           # 더 길다 → 위아래 자르기
        cut = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = pic.crop_bottom = cut
    round_corners(pic, w, h)


def round_corners(pic, w, h, r=0.25):
    pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {int(r / min(w, h) * 100000)}")


def movie(slide, name, x, y, w, h):
    mp4 = ROOT / "build" / "media" / f"{name}.mp4"
    png = ROOT / "build" / "media" / f"{name}.png"
    if not mp4.exists():
        missing.append(f"animation {name} (scripts/make_animations.py 실행 필요)")
        shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill="stage")
        return
    slide.shapes.add_movie(str(mp4), Inches(x), Inches(y), Inches(w), Inches(h),
                           poster_frame_image=str(png), mime_type="video/mp4")


# ── 레이아웃 ─────────────────────────────────────────────

def title(slide, s, color="ink"):
    text(slide, M, 0.5, W - 2 * M, 0.8, s, 30, color, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def footer(slide, deck, n, source=None):
    text(slide, M, 7.0, 4.5, 0.3, f"{COURSE['course']} · {deck['id']} {deck['title']}", 10, "muted")
    if source:                                       # 슬라이드 내용의 근거
        text(slide, W - M - 8.0, 7.0, 7.4, 0.3, f"출처: {source}", 10, "muted", align=PP_ALIGN.RIGHT)
    text(slide, W - M - 0.5, 7.0, 0.5, 0.3, str(n), 10, "muted", align=PP_ALIGN.RIGHT)


def lay_cover(slide, d, deck):
    background(slide, "stage")
    text(slide, 0.8, 2.0, 5.8, 0.4, f"{COURSE['course']}  ·  LECTURE {deck['id']}", 16, "amber", bold=True)
    text(slide, 0.8, 2.55, 5.8, 1.2, deck["title"], 54, "paper", bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.8, 3.95, 5.6, 1.0, deck.get("subtitle", ""), 20, SOFT_TEXT)
    text(slide, 0.8, 6.4, 5.6, 0.4, f"{COURSE['audience']}를 위한 첫걸음", 14, "muted")
    illustration(slide, d["illustration"], 6.9, 0.9, W - M - 6.9, 5.7, dark=True)


def lay_question(slide, d, deck):
    background(slide, "stage")
    text(slide, 0.8, 1.7, 5.6, 0.4, d.get("kicker", "질문"), 18, "amber", bold=True)
    text(slide, 0.8, 2.25, 5.6, 3.2, d["text"], 36, "paper", bold=True, spacing=1.3)
    illustration(slide, d["illustration"], 6.9, 0.9, W - M - 6.9, 5.7, dark=True)


def points_column(slide, points, x, y, w, gap=1.45):
    for i, p in enumerate(points):
        top = y + i * gap
        shape(slide, MSO_SHAPE.OVAL, x, top + 0.07, 0.26, 0.26, fill=p.get("color", "amber"))
        label_color = "ink" if p.get("color", "amber") == "amber" else p["color"]   # 흰 바탕 위 amber 글자는 흐림
        text(slide, x + 0.45, top, w - 0.45, 0.4, p["label"], 18, label_color, bold=True)
        text(slide, x + 0.45, top + 0.45, w - 0.45, gap - 0.5, p["text"], 15, "ink", spacing=1.25)


def lay_video(slide, d, deck):
    background(slide, "paper")
    title(slide, d["title"])
    if d.get("callout"):
        vw = 7.4
        movie(slide, d["animation"], M, 1.5, vw, vw * 9 / 16)
        cx = M + vw + 0.45
        points_column(slide, d.get("points", []), cx, 1.6, W - M - cx)
        box = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, M, 5.9, W - 2 * M, 0.9,
                    fill=CALLOUT_FILL, radius=0.15)
        box.name = "핵심 상자"
        text(slide, M + 0.3, 5.9, W - 2 * M - 0.6, 0.9, d["callout"], 15, "ink",
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    else:
        vw = 8.0
        movie(slide, d["animation"], M, 1.55, vw, vw * 9 / 16)
        cx = M + vw + 0.45
        points_column(slide, d.get("points", []), cx, 1.75, W - M - cx)


def lay_figure(slide, d, deck):
    background(slide, "paper")
    title(slide, d["title"])
    illustration(slide, d["illustration"], M, 1.55, 7.4, 5.0, dark=False)
    x = M + 7.4 + 0.5
    for i, step in enumerate(d["steps"]):
        top = 1.85 + i * 1.55
        badge(slide, x, top, 0.55, "amber", str(i + 1), 18, "stage")
        text(slide, x + 0.8, top + 0.05, W - M - x - 0.8, 1.3, step, 17, "ink", spacing=1.25)


def lay_summary(slide, d, deck):
    background(slide, "stage")
    title(slide, d["title"], "paper")
    gap = 0.35
    cw = (W - 2 * M - 2 * gap) / 3
    for i, c in enumerate(d["cards"]):
        x = M + i * (cw + gap)
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.8, cw, 4.6, fill="stage_2", radius=0.2)
        badge(slide, x + 0.4, 2.2, 0.8, c["color"], str(i + 1), 24, "paper")
        text(slide, x + 0.4, 3.3, cw - 0.8, 0.6, c["head"], 28, c["color"], bold=True)
        text(slide, x + 0.4, 4.05, cw - 0.8, 2.1, c["text"], 17, SOFT_TEXT, spacing=1.3)


def lay_quiz(slide, d, deck):
    background(slide, "paper")
    title(slide, d["title"])
    k, gap = len(d["items"]), 0.25
    rh = min(1.4, (5.15 - (k - 1) * gap) / k)         # 문항 수에 맞춰 줄 높이 조절
    dot = min(0.8, rh - 0.3)
    for i, it in enumerate(d["items"]):
        y = 1.6 + i * (rh + gap)
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, M, y, W - 2 * M, rh, fill="soft", radius=0.15)
        text(slide, M + 0.35, y, 0.8, rh, f"Q{i + 1}", 20, "amber", bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, M + 1.2, y, 8.6, rh, it["q"], 19, "ink", anchor=MSO_ANCHOR.MIDDLE, spacing=1.25)
        for j, mark in enumerate("OX"):
            b = shape(slide, MSO_SHAPE.OVAL, W - M - 2.15 + j * 1.0, y + (rh - dot) / 2, dot, dot,
                      fill="paper", line="line")
            tf = b.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = tf.paragraphs[0].add_run()
            style_run(r, 22, "proton" if mark == "X" else "electron", bold=True)
            r.text = mark
    d.setdefault("notes", "")
    d["notes"] += "\n[정답]\n" + "\n".join(
        f"Q{i + 1}. {it['a']} — {it['why']}" for i, it in enumerate(d["items"]))


def lay_objectives(slide, d, deck):
    background(slide, "paper")
    title(slide, d["title"])
    text(slide, M, 1.55, 6.6, 0.4, d.get("kicker", "이번 시간이 끝나면 이렇게 말할 수 있어요"), 16, "muted")
    colors = ("electron", "proton", "amber")
    for i, goal in enumerate(d["goals"]):
        y = 2.25 + i * 1.4
        badge(slide, M, y, 0.6, colors[i % 3], str(i + 1), 20, "paper")
        text(slide, M + 0.9, y - 0.05, 5.9, 1.2, goal, 20, "ink", anchor=MSO_ANCHOR.TOP, spacing=1.25)
    if d.get("poster"):                                # 이 강의 애니메이션의 한 장면
        png = ROOT / "build" / "media" / f"{d['poster']}.png"
        pw = W - M - 7.75
        pic = slide.shapes.add_picture(str(png), Inches(7.75), Inches(2.25), Inches(pw), Inches(pw * 9 / 16))
        round_corners(pic, pw, pw * 9 / 16)


def qr_image(url):
    out = ROOT / "build" / "figures" / f"qr_{hashlib.md5(url.encode()).hexdigest()[:8]}.png"
    if not out.exists():
        out.parent.mkdir(parents=True, exist_ok=True)
        qrcode.make(url, box_size=12, border=1).save(out)
    return out


def lay_try(slide, d, deck):
    background(slide, "stage")
    title(slide, d["title"], "paper")
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, M, 1.7, 3.6, 3.6, fill="paper", radius=0.2)
    slide.shapes.add_picture(str(qr_image(d["url"])), Inches(M + 0.3), Inches(2.0), Inches(3.0), Inches(3.0))
    text(slide, M, 5.45, 3.6, 0.6, d["url"], 10, SOFT_TEXT, align=PP_ALIGN.CENTER)
    x = M + 4.2
    text(slide, x, 1.7, W - M - x, 0.6, d["name"], 26, "amber", bold=True)
    for i, step in enumerate(d["steps"]):
        y = 2.6 + i * 1.05
        badge(slide, x, y, 0.5, "stage_2", str(i + 1), 16, "amber")
        text(slide, x + 0.75, y + 0.02, W - M - x - 0.75, 0.9, step, 18, "paper", spacing=1.2)


def lay_sources(slide, d, deck):
    background(slide, "paper")
    title(slide, d["title"])
    for i, it in enumerate(d["items"]):
        y = 1.6 + i * 1.02
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, M, y, W - 2 * M, 0.85, fill="soft", radius=0.12)
        text(slide, M + 0.3, y + 0.1, 9.4, 0.35, it["name"], 15, "ink", bold=True)
        text(slide, M + 0.3, y + 0.45, 9.4, 0.3, it["use"], 12, "muted")
        text(slide, W - M - 2.4, y, 2.1, 0.85, it["license"], 12, "muted",
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


LAYOUTS = {"objectives": lay_objectives, "try": lay_try, "sources": lay_sources, "cover": lay_cover, "question": lay_question, "video": lay_video,
           "figure": lay_figure, "summary": lay_summary, "quiz": lay_quiz}


def build(lecture):
    src = ROOT / "content" / lecture["dir"] / "slides.yaml"
    deck = yaml.safe_load(src.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for n, d in enumerate(deck["slides"], start=1):
        slide = prs.slides.add_slide(blank)
        LAYOUTS[d["layout"]](slide, d, deck)
        if d["layout"] not in ("cover",):
            footer(slide, deck, n, d.get("source"))
        if d.get("notes"):
            slide.notes_slide.notes_text_frame.text = d["notes"].strip()
    out = ROOT / "build" / "pptx" / f"{deck['id']}_{deck['title'].replace(' ', '')}.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(out)
    except PermissionError:                        # PowerPoint에서 열려 있으면 다른 이름으로
        out = out.with_stem(out.stem + "_새버전")
        prs.save(out)
        print("  (원본 파일이 열려 있어 '_새버전'으로 저장)")
    print(f"saved {out.relative_to(ROOT)}  ({len(deck['slides'])} slides)")


if __name__ == "__main__":
    wanted = set(sys.argv[1:])
    for lec in COURSE["lectures"]:
        if (not wanted or lec["id"] in wanted) and (ROOT / "content" / lec["dir"] / "slides.yaml").exists():
            build(lec)
    for m in dict.fromkeys(missing):
        print(f"  [자리표시] {m}")
